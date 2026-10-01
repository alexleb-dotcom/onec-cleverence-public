#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path
import argparse
import json
import re

DELIVERY_BINDING_PREFIX = "ONEC_RESULT_DELIVERY_BINDING_V1:"
DELIVERY_BINDING_FIELDS = (
    "task_id",
    "candidate_sha256",
    "package_binding_sha256",
    "change_items_sha256",
    "result_mode",
)
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit("DEPENDENCY_MISSING: install python-docx==1.2.0") from exc

ARTIFACTS = {
    "requirements": "Функциональная спецификация.docx",
    "manual_transfer": "Инструкция по внедрению.docx",
    "implementation_notes": "Особенности реализации.docx",
    "line_by_line": "Построчное обоснование изменений.docx",
}
PSEUDO_VERSION_RE = re.compile(r"(?:^|[_ .-])(v\d+|final\d*|fix\d*|new)(?:[_ .-]|$)", re.I)
INVALID_FILENAME_RE = re.compile(r'[<>:"/\\\\|?*]')


def _require(payload: dict, *keys: str) -> None:
    missing = [key for key in keys if payload.get(key) in (None, "", [], {})]
    if missing:
        raise ValueError("missing required fields: " + ", ".join(missing))


def _text(value) -> str:
    return "" if value is None else str(value)


def _safe_filename(name: str) -> str:
    if not name.lower().endswith(".docx"):
        raise ValueError("filename must end with .docx")
    if INVALID_FILENAME_RE.search(name) or name in {".", ".."}:
        raise ValueError("filename contains forbidden characters")
    if PSEUDO_VERSION_RE.search(name[:-5]):
        raise ValueError("pseudo-version filename suffix is forbidden")
    return name


def canonical_filename(artifact: str, payload: dict) -> str:
    requested = payload.get("filename")
    return _safe_filename(requested) if requested else ARTIFACTS[artifact]


def _delivery_binding_token(artifact: str, payload: dict) -> str | None:
    binding = payload.get("delivery_binding")
    if binding in (None, {}):
        return None
    if not isinstance(binding, dict):
        raise ValueError("delivery_binding must be an object")
    missing = [field for field in DELIVERY_BINDING_FIELDS if binding.get(field) in (None, "")]
    if missing:
        raise ValueError("delivery_binding missing fields: " + ", ".join(missing))
    for field in ("candidate_sha256", "package_binding_sha256", "change_items_sha256"):
        value = str(binding.get(field) or "")
        if not SHA256_RE.fullmatch(value):
            raise ValueError(f"delivery_binding {field} must be lowercase sha256")
    normalized = {field: str(binding[field]) for field in DELIVERY_BINDING_FIELDS}
    normalized["artifact_kind"] = artifact
    return DELIVERY_BINDING_PREFIX + json.dumps(
        normalized,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _base(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(1.5)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(size)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    doc.core_properties.title = title
    return doc


def _meta(doc: Document, rows: list[tuple[str, str]]) -> None:
    for label, value in rows:
        p = doc.add_paragraph()
        r = p.add_run(label + ": ")
        r.bold = True
        p.add_run(_text(value))


def _heading(doc: Document, value: str, level: int = 1) -> None:
    doc.add_heading(value, level=level)


def _body(doc: Document, value) -> None:
    if value in (None, "", [], {}):
        return
    if isinstance(value, list):
        for item in value:
            doc.add_paragraph(_text(item), style="List Bullet")
        return
    for block in _text(value).split("\n"):
        doc.add_paragraph(block)


def _code(doc: Document, value: str) -> None:
    p = doc.add_paragraph()
    for line in _text(value).splitlines() or [""]:
        r = p.add_run(line + "\n")
        r.font.name = "Courier New"
        r.font.size = Pt(9)


def _shade(cell, fill: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    header_props = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        cell.text = label
        _shade(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(9)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = _text(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)


def _rationale(doc: Document, rationale: str, standard: str | None) -> None:
    p = doc.add_paragraph()
    p.add_run("Обоснование: ").bold = True
    p.add_run(_text(rationale))
    if standard:
        p = doc.add_paragraph()
        p.add_run("Применимый стандарт/правило: ").bold = True
        p.add_run(_text(standard))


def _requirements(payload: dict) -> Document:
    _require(payload, "project", "task", "need_problem", "target_outcome",
             "functional_rules", "acceptance", "proof_boundary", "requirements_status")
    if payload["requirements_status"] == "REQUIREMENTS_BLOCKED" and not payload.get("open_questions"):
        raise ValueError("blocked requirements must expose open_questions")
    doc = _base(payload.get("title") or "Функциональная спецификация")
    _meta(doc, [("Проект", payload["project"]), ("Задача", payload["task"])])
    for title, key in (
        ("Потребность / проблема", "need_problem"),
        ("Целевой результат", "target_outcome"),
    ):
        _heading(doc, title); _body(doc, payload[key])
    _heading(doc, "Границы")
    _heading(doc, "Входит", 2); _body(doc, payload.get("scope_in") or ["Не указано"])
    _heading(doc, "Не входит", 2); _body(doc, payload.get("scope_out") or ["Не указано"])
    _heading(doc, "Функциональное поведение и материальные правила")
    _body(doc, payload["functional_rules"])
    for title, key in (
        ("Источник истины / идентичность / данные", "source_identity_data"),
        ("Процесс / состояния / повтор / ошибки", "process_states_errors"),
        ("Интеграционный контракт", "integration_contract"),
    ):
        if payload.get(key):
            _heading(doc, title); _body(doc, payload[key])
    _heading(doc, "Приемка")
    rows = [[r.get("case",""), r.get("preconditions",""), r.get("action",""),
             r.get("expected",""), r.get("oracle","")] for r in payload["acceptance"]]
    _table(doc, ["Случай","Предусловия","Действие","Ожидаемый результат","Оракул"], rows)
    for title, key in (
        ("Открытые вопросы / блокеры", "open_questions"),
        ("Допущения", "assumptions"),
        ("Предлагаемое решение", "proposed_solution"),
    ):
        if payload.get(key):
            _heading(doc, title); _body(doc, payload[key])
    _heading(doc, "Основания и граница доказанности"); _body(doc, payload["proof_boundary"])
    labels = {
        "REQUIREMENTS_READY": "Требования готовы",
        "REQUIREMENTS_READY_WITH_ASSUMPTIONS": "Требования готовы с допущениями",
        "REQUIREMENTS_BLOCKED": "Требования заблокированы",
    }
    state = payload["requirements_status"]
    if state not in labels:
        raise ValueError("unsupported requirements_status")
    p = doc.add_paragraph(); p.add_run("Статус требований: ").bold = True
    p.add_run(f"{labels[state]} ({state})")
    return doc


def _manual(payload: dict) -> Document:
    _require(
        payload,
        "project",
        "task",
        "target_identity",
        "purpose",
        "implementation_boundaries",
        "deployment_sequence",
        "verification_matrix",
        "final_control_checklist",
        "changed_object_map",
        "proof_boundary",
    )
    doc = _base(payload.get("title") or "Инструкция по внедрению")
    _meta(doc, [("Проект", payload["project"]), ("Задача", payload["task"]),
                ("Целевая база / артефакт", payload["target_identity"])])

    _heading(doc, "Назначение и границы")
    _body(doc, payload["purpose"])
    _body(doc, payload["implementation_boundaries"])
    if payload.get("do_not_change"):
        _heading(doc, "Не изменять / не делать")
        _body(doc, payload["do_not_change"])

    _heading(doc, "Создаваемые объекты")
    created = payload.get("created_objects") or []
    if not created:
        doc.add_paragraph("Нет")
    for item in created:
        _require(item, "object", "properties", "rationale")
        _heading(doc, "Объект: " + item["object"], 2)
        _table(doc, ["Свойство","Значение"],
               [[p.get("name",""), p.get("value","")] for p in item["properties"]])
        _rationale(doc, item["rationale"], item.get("standard_rule"))

    _heading(doc, "Изменяемые объекты")
    modified = payload.get("modified_objects") or []
    if not modified:
        doc.add_paragraph("Нет")
    for item in modified:
        _require(item, "object", "properties", "rationale")
        _heading(doc, "Объект: " + item["object"], 2)
        _table(doc, ["Свойство","Было","Стало"],
               [[p.get("name",""), p.get("before",""), p.get("after","")] for p in item["properties"]])
        _rationale(doc, item["rationale"], item.get("standard_rule"))

    _heading(doc, "Код")
    changes = payload.get("code_changes") or []
    if not changes:
        doc.add_paragraph("Нет")
    for item in changes:
        _require(item, "object", "member", "anchor", "before", "after", "rationale")
        _heading(doc, "Объект: " + item["object"], 2)
        _meta(doc, [("Изменения", item["member"]), ("Якорь / место изменения", item["anchor"])])
        if item.get("dependencies"):
            _meta(doc, [("Зависимости / порядок", "; ".join(map(_text, item["dependencies"])))])
        _heading(doc, "Было", 2)
        _code(doc, item["before"])
        _heading(doc, "Стало", 2)
        _code(doc, item["after"])
        _rationale(doc, item["rationale"], item.get("standard_rule"))
        if item.get("explanation"):
            p = doc.add_paragraph()
            p.add_run("Пояснение: ").bold = True
            p.add_run(_text(item["explanation"]))

    if payload.get("preconditions"):
        _heading(doc, "Предусловия")
        _body(doc, payload["preconditions"])
    if payload.get("migration"):
        _heading(doc, "Миграция / инициализация / одноразовые действия")
        _body(doc, payload["migration"])

    _heading(doc, "Порядок внедрения")
    _body(doc, payload["deployment_sequence"])

    if payload.get("static_verification"):
        _heading(doc, "Статическая проверка после внедрения")
        _body(doc, payload["static_verification"])

    _heading(doc, "Матрица проверки")
    verification_rows = []
    for row in payload["verification_matrix"]:
        _require(row, "action", "expected")
        verification_rows.append([row["action"], row["expected"]])
    _table(doc, ["Действие / сценарий", "Ожидаемый результат"], verification_rows)

    if payload.get("runtime_verification"):
        _heading(doc, "Проверка выполнения")
        _body(doc, payload["runtime_verification"])

    _heading(doc, "Финальный статический контроль")
    _body(doc, payload["final_control_checklist"])

    _heading(doc, "Карта изменённых объектов")
    changed_rows = []
    for row in payload["changed_object_map"]:
        _require(row, "object", "change")
        changed_rows.append([row["object"], row["change"]])
    _table(doc, ["Объект", "Изменение"], changed_rows)

    if payload.get("blocking_choices"):
        _heading(doc, "Нерешённые выборы / блокеры")
        _body(doc, payload["blocking_choices"])

    _heading(doc, "Граница доказанности")
    _body(doc, payload["proof_boundary"])
    return doc

def _implementation_notes(payload: dict) -> Document:
    _require(payload, "tz_number", "project", "task", "platform_version",
             "configuration_name_version", "rows")
    doc = _base("Особенности реализации")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    _meta(doc, [
        ("Номер ТЗ", payload["tz_number"]),
        ("Проект", payload["project"]),
        ("Задача", payload["task"]),
        ("Версия платформы", payload["platform_version"]),
        ("Наименование и версия конфигурации", payload["configuration_name_version"]),
    ])
    rows = []
    for row in payload["rows"]:
        _require(row, "container", "configuration_object", "status", "description", "standard_rule")
        description = row["description"] + " Стандарт/правило: " + row["standard_rule"]
        rows.append([row["container"], row["configuration_object"],
                     row.get("procedure_function") or "—", row["status"], description])
    _table(doc, ["Контейнер","ОбъектКонфигурации","Процедура/Функция","Статус","ОписаниеИзменений"], rows)
    return doc


def _line_by_line(payload: dict) -> Document:
    _require(payload, "explicit_user_request", "project", "task", "purpose", "objects")
    if payload["explicit_user_request"] is not True:
        raise ValueError("line_by_line artifact requires explicit_user_request=true")
    doc = _base("Построчное обоснование изменений")
    meta = []
    if payload.get("tz_number"):
        meta.append(("Номер ТЗ", payload["tz_number"]))
    meta.extend([("Проект", payload["project"]), ("Задача", payload["task"])])
    _meta(doc, meta)
    _heading(doc, "Назначение документа"); _body(doc, payload["purpose"])
    _heading(doc, "Правила чтения")
    _body(doc, payload.get("reading_rules") or [
        "Строки с '-' удалены.",
        "Строки с '+' добавлены.",
        "Строки без маркера приведены как контекст.",
    ])
    for obj in payload["objects"]:
        _require(obj, "object", "changes")
        _heading(doc, "Объект: " + obj["object"])
        for change in obj["changes"]:
            _require(change, "member", "status", "reason", "what_changed", "behavior_impact", "diff")
            _heading(doc, change["member"], 2)
            _meta(doc, [("Статус", change["status"])])
            if change.get("task_relation"):
                _meta(doc, [("Связь с ТЗ/задачей", change["task_relation"])])
            if change.get("location"):
                _meta(doc, [("Точное место изменения", change["location"])])
            _heading(doc, "Причина изменения", 2); _body(doc, change["reason"])
            _heading(doc, "Что изменено", 2); _body(doc, change["what_changed"])
            _heading(doc, "Влияние на поведение", 2); _body(doc, change["behavior_impact"])
            _heading(doc, "Точный diff", 2); _code(doc, change["diff"])
    return doc


RENDERERS = {
    "requirements": _requirements,
    "manual_transfer": _manual,
    "implementation_notes": _implementation_notes,
    "line_by_line": _line_by_line,
}


def render_artifact(artifact: str, payload: dict, output_dir: Path) -> Path:
    if artifact not in RENDERERS:
        raise ValueError("unknown artifact: " + artifact)
    output_dir.mkdir(parents=True, exist_ok=True)
    output = (output_dir / canonical_filename(artifact, payload)).resolve()
    if output.parent != output_dir.resolve():
        raise ValueError("output filename escapes output directory")
    doc = RENDERERS[artifact](payload)
    doc.core_properties.subject = "User-facing projection; canonical requirements/proof owners remain external"
    binding_token = _delivery_binding_token(artifact, payload)
    if binding_token is not None:
        # Invisible file-bound identity used only by the existing Result Delivery
        # final-set verifier. It does not create semantic or release proof.
        doc.core_properties.identifier = binding_token
    doc.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="Render one owner-approved user-facing artifact as a separate DOCX file.")
    parser.add_argument("--artifact", required=True, choices=sorted(RENDERERS))
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    try:
        output = render_artifact(args.artifact, payload, Path(args.output_dir))
    except (ValueError, TypeError) as exc:
        print(json.dumps({"result":"FAIL","artifact":args.artifact,"error":str(exc)}, ensure_ascii=False, indent=2))
        return 2
    print(json.dumps({"result":"PASS","artifact":args.artifact,"output":str(output),"filename":output.name}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
